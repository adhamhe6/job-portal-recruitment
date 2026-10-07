"""Realistic résumé fixtures generated at test time (reportlab PDFs, python-docx DOCX files, hostile variants).

Three distinct, fictional profiles — a backend engineer, a frontend engineer and a nurse — each available as text, PDF
and DOCX. Nothing here is real personal data.
"""

from __future__ import annotations

import io
import textwrap
import zipfile
from collections.abc import Sequence

BACKEND_TEXT = """Jane Doe
Senior Backend Engineer | Python | AWS
Berlin, Germany | jane.doe@example.com | +49 151 2345 6789 | linkedin.com/in/janedoe | github.com/janedoe

PROFESSIONAL SUMMARY
Backend engineer with 8+ years of experience building scalable APIs and data pipelines. Passionate about reliability.

SKILLS
Languages: Python, Go, SQL, Bash
Frameworks: FastAPI, Django, Flask
Databases: PostgreSQL, Redis, MongoDB
DevOps: Docker, Kubernetes, Terraform, CI/CD, GitHub Actions, AWS

WORK EXPERIENCE
Senior Backend Engineer
Acme Corp, Berlin | Jan 2020 – Present
• Designed REST APIs with FastAPI serving 2M requests per day
• Led migration from MySQL to PostgreSQL; used Go for a small tool
• Mentored 4 junior engineers

Backend Developer — Initech GmbH                         03/2016 - 12/2019
• Built Django services and Celery workers
• Introduced Docker and CI/CD pipelines

Software Engineer at Globex (2013-2016)
Developed internal tools in Java and SQL.

EDUCATION
Technical University of Berlin
B.Sc. in Computer Science, 2009 – 2013

CERTIFICATIONS
AWS Certified Solutions Architect – Associate (Amazon Web Services), 2021
Certified Kubernetes Administrator (CKA)

LANGUAGES
English (Fluent), German (Native), French - Basic
"""

FRONTEND_TEXT = """Alex Kim
alex.kim@mail.example | 415-555-0199 | San Francisco, CA | github.com/alexkim

Profile
Frontend developer focused on accessible, performant user interfaces.

Technical Skills
JavaScript, TypeScript, React, Next.js, Redux, HTML5, CSS3, Tailwind CSS, Jest, Cypress, Figma

Experience
Frontend Developer, Globex Corporation                      2021 - Present
Built a design system in React and TypeScript used by 12 teams.
Improved Lighthouse performance scores from 62 to 95.
UI Developer, Initech Labs                                  Jun 2018 - Dec 2020
Maintained marketing sites in HTML, CSS and jQuery.

Education
Stanford University
BS in Computer Science (2014 - 2018)
"""

NURSE_TEXT = """MARIA GONZALEZ, RN
Registered Nurse
Austin, TX | (512) 555-0142 | maria.gonzalez@example.com

SUMMARY
Compassionate registered nurse with 6 years of experience in critical care and patient assessment.

CLINICAL EXPERIENCE
ICU Registered Nurse
St. David's Medical Center - Austin, TX
March 2019 - Present
• Monitored patient vitals and administered medication (5 ml doses) per physician orders
• Provided Critical Care to ventilated patients; maintained Electronic Health Records in Epic

Staff Nurse
Seton Hospital, Austin, TX
06/2015 - 02/2019
• Delivered patient care on a 30-bed medical-surgical unit

EDUCATION
BSN, Nursing - University of Texas at Austin, 2011 - 2015

LICENSES & CERTIFICATIONS
Registered Nurse (RN) License - State of Texas
BLS Certified - American Heart Association, 2023
ACLS Certified, 2022

LANGUAGES
English (Native), Spanish (Fluent)

INTERESTS
Hiking, photography, reading about React
"""

HEADINGS = {
    "PROFESSIONAL SUMMARY",
    "SKILLS",
    "WORK EXPERIENCE",
    "EDUCATION",
    "CERTIFICATIONS",
    "LANGUAGES",
    "SUMMARY",
    "CLINICAL EXPERIENCE",
    "LICENSES & CERTIFICATIONS",
    "INTERESTS",
    "Profile",
    "Technical Skills",
    "Experience",
    "Education",
}


def make_pdf(text: str, *, wrap: int = 98) -> bytes:
    """A text-based PDF: headings in bold, one text object per line, simple pagination."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    _, height = A4
    left, top, bottom, leading = 50, height - 60, 55, 14
    y = top
    for raw in text.splitlines():
        lines = textwrap.wrap(raw, wrap) or [""]
        for line in lines:
            if y < bottom:
                c.showPage()
                y = top
            c.setFont(
                "Helvetica-Bold" if raw.strip() in HEADINGS else "Helvetica",
                11 if raw.strip() in HEADINGS else 10,
            )
            if line:
                c.drawString(left, y, line)
            y -= leading
    c.save()
    return buf.getvalue()


def make_docx(text: str, *, tables: bool = False) -> bytes:
    """A Word document: headings as Heading 1, bullets as List Bullet; optionally the skills block inside a table."""
    from docx import Document

    doc = Document()
    skill_lines: list[str] = []
    in_skills = False
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped in HEADINGS:
            in_skills = stripped.lower() in ("skills", "technical skills")
            doc.add_heading(stripped, level=1)
            continue
        if tables and in_skills and stripped:
            skill_lines.append(stripped)
            continue
        if skill_lines and tables and not in_skills:
            _skills_table(doc, skill_lines)
            skill_lines = []
        if stripped.startswith("• "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(stripped)
    if skill_lines:
        _skills_table(doc, skill_lines)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _skills_table(doc: object, lines: Sequence[str]) -> None:
    table = doc.add_table(rows=len(lines), cols=2)  # type: ignore[attr-defined]
    for row, line in zip(table.rows, lines, strict=True):
        label, _, rest = line.partition(":")
        row.cells[0].text = label.strip()
        row.cells[1].text = rest.strip() or label.strip()


def docx_with_header(text: str, header: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = header
    for line in text.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def backend_pdf() -> bytes:
    return make_pdf(BACKEND_TEXT)


def backend_docx() -> bytes:
    return make_docx(BACKEND_TEXT)


def frontend_pdf() -> bytes:
    return make_pdf(FRONTEND_TEXT)


def frontend_docx(tables: bool = False) -> bytes:
    return make_docx(FRONTEND_TEXT, tables=tables)


def nurse_pdf() -> bytes:
    return make_pdf(NURSE_TEXT)


def nurse_docx() -> bytes:
    return make_docx(NURSE_TEXT)


# --- hostile / broken inputs -----------------------------------------------------------------------------------------


def encrypted_pdf(password: str = "s3cret-pw") -> bytes:  # noqa: S107
    """A PDF that needs a password to open (RC4, so no crypto backend is required to create or to reject it)."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(backend_pdf())))
    writer.encrypt(user_password=password, owner_password=password, algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def malformed_pdf() -> bytes:
    return (
        b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog /Pages 99 0 R >>\nendobj\n"
        + b"\x00garbage\xff" * 200
        + b"\n%%EOF"
    )


def truncated_pdf() -> bytes:
    data = backend_pdf()
    return data[: len(data) // 3]


def scanned_pdf() -> bytes:
    """A PDF whose only content is a picture (no text layer): what a scan looks like to a text extractor."""
    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (400, 300), "white")
    for x in range(40, 360, 4):
        for y in range(40, 60):
            img.putpixel((x, y), (0, 0, 0))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.drawImage(ImageReader(img), 50, 400, width=400, height=300)
    c.save()
    return buf.getvalue()


def blank_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.showPage()
    c.save()
    return buf.getvalue()


def empty_docx() -> bytes:
    from docx import Document

    buf = io.BytesIO()
    Document().save(buf)
    return buf.getvalue()


CONTENT_TYPES_XML = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    b'<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    b"</Types>"
)


def zip_bytes(members: dict[str, bytes], *, compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def zip_bomb_docx(uncompressed_mb: int = 60) -> bytes:
    """A tiny DOCX-shaped archive that expands to ``uncompressed_mb`` megabytes."""
    chunk = b"<w:p>" + b"A" * 1000 + b"</w:p>"
    body = (
        b'<w:document xmlns:w="x"><w:body>'
        + chunk * (uncompressed_mb * 1024 * 1024 // len(chunk))
        + b"</w:body></w:document>"
    )
    return zip_bytes({"[Content_Types].xml": CONTENT_TYPES_XML, "word/document.xml": body})


def traversal_docx() -> bytes:
    return zip_bytes(
        {
            "[Content_Types].xml": CONTENT_TYPES_XML,
            "word/document.xml": b"<w:document/>",
            "../../evil.txt": b"x",
        }
    )


def macro_docx() -> bytes:
    return zip_bytes(
        {
            "[Content_Types].xml": CONTENT_TYPES_XML,
            "word/document.xml": b"<w:document/>",
            "word/vbaProject.bin": b"\x00" * 10,
        }
    )


def xlsx_like() -> bytes:
    """A valid ZIP/OOXML package that is *not* a Word document."""
    return zip_bytes(
        {
            "[Content_Types].xml": b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
            "xl/workbook.xml": b"<workbook/>",
        }
    )


OLE_DOC = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 600  # legacy .doc container
EXE = b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 200
