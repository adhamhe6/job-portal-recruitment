"""Text extraction from PDF and DOCX, including the failure modes."""

from __future__ import annotations

import io
import time

import pytest

from app.resume import extract as ex
from app.resume.extract import ExtractionError, extract_document, extract_sync, normalize_text
from app.resume.validation import DocumentKind
from tests import fixtures_resumes as fx

PDF, DOCX = DocumentKind.PDF, DocumentKind.DOCX
MAX = 200_000


def test_pdf_text_is_extracted_line_by_line():
    r = extract_sync(fx.backend_pdf(), PDF, max_chars=MAX)
    assert r.page_count == 1 and not r.was_truncated and r.pages_failed == 0
    assert "Jane Doe" in r.text and "jane.doe@example.com" in r.text
    assert "PROFESSIONAL SUMMARY" in r.text.splitlines()
    assert "Python, Go, SQL, Bash" in r.text
    assert "03/2016 - 12/2019" in r.text


def test_multi_page_pdf_and_page_cap():
    long_text = "\n".join(f"Line {i} of a very long résumé about Python engineering experience" for i in range(2500))
    data = fx.make_pdf(long_text)
    full = extract_sync(data, PDF, max_chars=10_000_000, max_pages=1000)
    assert full.page_count > 30 and "Line 2499" in full.text and not full.was_truncated
    capped = extract_sync(data, PDF, max_chars=10_000_000, max_pages=30)
    assert capped.page_count == full.page_count  # reports the real page count
    assert capped.was_truncated and "Line 2499" not in capped.text and "Line 0 " in capped.text


def test_text_truncated_to_the_character_limit():
    r = extract_sync(fx.backend_pdf(), PDF, max_chars=300)
    assert r.was_truncated and len(r.text) <= 300
    assert r.text.startswith("Jane Doe")


def test_docx_paragraphs_and_headings():
    r = extract_sync(fx.backend_docx(), DOCX, max_chars=MAX)
    assert r.page_count is None
    assert "Jane Doe" in r.text and "WORK EXPERIENCE" in r.text.splitlines()
    assert "Designed REST APIs with FastAPI" in r.text  # bullet paragraphs (List Bullet style)


def test_docx_tables_are_included_in_document_order():
    r = extract_sync(fx.frontend_docx(tables=True), DOCX, max_chars=MAX)
    assert "JavaScript, TypeScript, React" in r.text
    lines = r.text.splitlines()
    assert lines.index("Technical Skills") < lines.index("Experience")


def test_docx_merged_table_cells_are_not_repeated():
    from docx import Document

    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    merged = t.cell(0, 0).merge(t.cell(0, 1))
    merged.text = "Merged cell with enough characters"
    t.cell(0, 2).text = "Right cell"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_sync(buf.getvalue(), DOCX, max_chars=MAX).text
    assert text.count("Merged cell with enough characters") == 1 and "Right cell" in text


def test_docx_header_text_is_included():
    data = fx.docx_with_header("Experience\nEngineer at Acme Corp 2020 - 2022\nBuilt things for customers", "Jordan Lee | jordan@example.com")
    assert "jordan@example.com" in extract_sync(data, DOCX, max_chars=MAX).text.splitlines()[0]


def test_docx_hyperlink_targets_are_returned_as_links():
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document()
    doc.add_paragraph("LinkedIn profile for a reasonably long résumé body text goes here")
    doc.part.relate_to("https://www.linkedin.com/in/jordan", RT.HYPERLINK, is_external=True)
    buf = io.BytesIO()
    doc.save(buf)
    assert "https://www.linkedin.com/in/jordan" in extract_sync(buf.getvalue(), DOCX, max_chars=MAX).links


def test_encrypted_pdf_is_rejected_gracefully():
    with pytest.raises(ExtractionError) as exc:
        extract_sync(fx.encrypted_pdf(), PDF, max_chars=MAX)
    assert exc.value.code == "ENCRYPTED_DOCUMENT" and "password" in exc.value.message.lower()


@pytest.mark.parametrize("data", [fx.malformed_pdf(), b"%PDF-1.4\nthis is not really a pdf", b"%PDF-", fx.truncated_pdf()])
def test_malformed_pdf_is_a_clean_failure(data):
    with pytest.raises(ExtractionError) as exc:
        extract_sync(data, PDF, max_chars=MAX)
    assert exc.value.code in {"MALFORMED_DOCUMENT", "NO_TEXT_EXTRACTED"}


def test_malformed_docx_is_a_clean_failure():
    for data in (b"PK\x03\x04 not a zip", fx.zip_bytes({"hello.txt": b"hi"}), fx.zip_bytes({"[Content_Types].xml": fx.CONTENT_TYPES_XML, "word/document.xml": b"<broken"})):
        with pytest.raises(ExtractionError) as exc:
            extract_sync(data, DOCX, max_chars=MAX)
        assert exc.value.code == "MALFORMED_DOCUMENT"


@pytest.mark.parametrize(("data", "kind"), [(fx.scanned_pdf(), PDF), (fx.blank_pdf(), PDF), (fx.empty_docx(), DOCX)])
def test_scanned_or_empty_documents_report_no_text_honestly(data, kind):
    with pytest.raises(ExtractionError) as exc:
        extract_sync(data, kind, max_chars=MAX)
    assert exc.value.code == "NO_TEXT_EXTRACTED"
    assert "scanned" in exc.value.message and "OCR is not enabled" in exc.value.message


def test_error_messages_never_contain_document_text():
    secret = "Jane Doe jane.doe@example.com"
    for data, kind in ((fx.malformed_pdf(), PDF), (fx.encrypted_pdf(), PDF), (fx.scanned_pdf(), PDF), (fx.truncated_pdf(), PDF)):
        with pytest.raises(ExtractionError) as exc:
            extract_sync(data, kind, max_chars=MAX)
        assert secret not in exc.value.message and "jane" not in exc.value.message.lower()


async def test_extract_document_async_wrapper_and_unexpected_errors_are_converted(monkeypatch):
    r = await extract_document(fx.backend_pdf(), PDF, max_chars=MAX)
    assert "Jane Doe" in r.text

    def boom(*a, **k):
        raise RecursionError("jane.doe@example.com leaked")

    monkeypatch.setattr(ex, "extract_sync", boom)
    with pytest.raises(ExtractionError) as exc:
        await extract_document(b"x", PDF, max_chars=MAX)
    assert exc.value.code == "MALFORMED_DOCUMENT" and "leaked" not in exc.value.message


async def test_extraction_timeout(monkeypatch):
    monkeypatch.setattr(ex, "extract_sync", lambda *a, **k: time.sleep(1.5))
    with pytest.raises(ExtractionError) as exc:
        await extract_document(b"x", PDF, max_chars=MAX, timeout_seconds=0.1)
    assert exc.value.code == "EXTRACTION_TIMEOUT"


# --- normalisation -------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("develop-\nment of apps", "development of apps"),
        ("2018-\n2021", "2018-\n2021"),  # numeric ranges are not hyphenation
        ("full-\nStack", "full-\nStack"),  # only joins lower-case halves
        ("a   b\t\tc", "a b c"),
        ("line1\r\nline2\rline3", "line1\nline2\nline3"),
        ("a\n\n\n\n\nb", "a\n\nb"),
        ("  padded  \n  lines  ", "padded\nlines"),
        ("ﬁnal ﬂow", "final flow"),  # ligatures folded (NFKC)
        ("zero​width and­soft", "zero width andsoft"),
        ("nul\x00byte", "nul byte"),
        ("nbsp here", "nbsp here"),
        ("", ""),
    ],
)
def test_normalize_text(raw, expected):
    assert normalize_text(raw) == expected
